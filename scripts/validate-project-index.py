"""Reconcile public project fields; never read participant IDs or private links."""
import argparse, hashlib, json, re, subprocess
from pathlib import Path
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('workbook', type=Path)
parser.add_argument('--output', type=Path, default=Path('docs/project-index-validation-2026-10-13.json'))
args = parser.parse_args()
node = subprocess.check_output(['node', '--import', 'tsx', '--input-type=module', '-e',
    'const {researchProjectRegister:r}=await import("./src/data/researchInitiatives.ts");console.log(JSON.stringify(r));'], text=True)
projects = json.loads(node)
sheet = openpyxl.load_workbook(args.workbook, data_only=True)['Sheet1']
rows = [item for project in projects for item in [project, *project.get('subProjects', [])]]
def counts(value):
    if isinstance(value, (int, float)):
        return [int(value)]
    text = re.sub(r'\([^)]*\)', '', str(value or '')).replace(',', '')
    return [int(n) for n in re.findall(r'\d+', text)]
errors = []
for item in rows:
    row = item['sourceRow']
    if not sheet.cell(row, 2).value:
        errors.append({'row': row, 'field': 'title', 'error': 'source row missing'})
    for field, column in [('targetedParticipants', 8), ('currentParticipants', 9)]:
        if counts(item.get(field)) != counts(sheet.cell(row, column).value):
            errors.append({'row': row, 'field': field, 'error': 'numeric count differs'})
covered = sorted(item['sourceRow'] for item in rows)
expected = [r for r in range(2, 80) if sheet.cell(r, 2).value]
if covered != expected:
    errors.append({'field': 'coverage', 'error': 'missing or duplicated source rows'})
report = {
    'checkedOn': '2026-10-13', 'sourceWorkbook': args.workbook.name,
    'sourceSha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
    'sheet': 'Sheet1', 'sourceRows': '2–79',
    'projects': len(projects), 'subprojects': len(rows)-len(projects),
    'coverageChecks': len(rows), 'countComparisons': 2*len(rows), 'errors': errors,
    'result': 'passed' if not errors else 'failed',
    'scope': 'Source reconciliation only. Current participant figures remain withheld from the UI pending Statistics Unit publication approval. No patient-level dataset, identifiers or private source links were read.'
}
args.output.parent.mkdir(parents=True, exist_ok=True)
args.output.write_text(json.dumps(report, indent=2)+'\n', encoding='utf8')
print(json.dumps(report))
raise SystemExit(bool(errors))
